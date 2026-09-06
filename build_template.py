"""
Builds template_lhp.docx from scratch.
Placeholders use {{TOKEN}} syntax and each lives in its own dedicated run,
so app.py can do a single whole-string replace per run with zero ambiguity
(this avoids the substring-ordering bugs from earlier versions of this project
 -- see CLAUDE.md).
"""
import docx
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial Narrow"

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge, {'sz': 4, 'val': 'single', 'color': '000000'})
        tag = 'w:{}'.format(edge)
        element = OxmlElement(tag)
        for key, val in edge_data.items():
            element.set(qn('w:{}'.format(key)), str(val))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def add_run(p, text, bold=False, underline=False, italic=False, size=12, font=FONT):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.underline = underline
    r.italic = italic
    rpr = r._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font)
    return r


def new_para(doc_or_cell, align=None, spacing_after=0):
    p = doc_or_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(spacing_after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    return p


def build():
    doc = Document()

    # ---- page setup ----
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(33.02)  # Legal/Folio-ish, matches Indonesian gov docs (F4)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(12)

    # ---- KOP ----
    p = new_para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, "RESIMEN TARUNA DAN SISWA", size=12)

    p = new_para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, "{{KOP_BATALYON}}", size=12)

    p = new_para(doc, WD_ALIGN_PARAGRAPH.CENTER, spacing_after=6)
    add_run(p, "{{JUDUL_LAPORAN}}", bold=True, underline=True, size=13)

    # ---- Table 0: identitas ----
    t0 = doc.add_table(rows=2, cols=3)
    t0.alignment = WD_TABLE_ALIGNMENT.CENTER
    t0.style = 'Table Grid'
    widths = [Inches(1.6), Inches(2.6), Inches(1.9)]
    for row in t0.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]

    # row0: blank | logo | ACCEPTED/INPUT
    c00, c01, c02 = t0.rows[0].cells
    p = c00.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = c01.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture('static/logo_akpol.png', height=Cm(2.4))

    p = c02.paragraphs[0]
    add_run(p, "ACCEPTED:", bold=True, size=11)
    p2 = c02.add_paragraph()
    add_run(p2, "INPUT :", bold=True, size=11)

    # row1: NO AKADEMI/data | NAMA TARUNA/data | KOMPI x / PLETON y
    c10, c11, c12 = t0.rows[1].cells
    p = c10.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "NO AKADEMI", bold=True, size=11)
    p2 = c10.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p2, "{{NO_AKADEMI}}", bold=True, italic=True, size=11)

    p = c11.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "NAMA TARUNA", bold=True, size=11)
    p2 = c11.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p2, "{{NAMA_TARUNA}}", bold=True, italic=True, size=11)

    p = c12.paragraphs[0]
    add_run(p, "{{KOMPI_LABEL}}", bold=True, size=11)
    p2 = c12.add_paragraph()
    add_run(p2, "{{PELETON_LABEL}}", bold=True, size=11)

    new_para(doc, spacing_after=4)

    # ---- Table 1: kegiatan / uraian / disposisi ----
    t1 = doc.add_table(rows=5, cols=1)
    t1.style = 'Table Grid'
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in t1.rows:
        row.cells[0].width = Inches(6.1)

    r0 = t1.rows[0].cells[0]
    p = r0.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Inches(2.2))
    add_run(p, "NAMA KEGIATAN\t             :   ", bold=True, size=11)
    add_run(p, "{{NAMA_KEGIATAN}}", bold=True, size=11)
    p2 = r0.add_paragraph()
    p2.paragraph_format.tab_stops.add_tab_stop(Inches(2.2))
    add_run(p2, "WAKTU PELAKSANAAN          :    ", bold=True, size=11)
    add_run(p2, "{{WAKTU_PELAKSANAAN}}", bold=True, size=11)
    p3 = r0.add_paragraph()
    p3.paragraph_format.tab_stops.add_tab_stop(Inches(2.2))
    add_run(p3, "TEMPAT\t\t             :    ", bold=True, size=11)
    add_run(p3, "{{TEMPAT}}", bold=True, size=11)

    r1 = t1.rows[1].cells[0]
    p = r1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "URAIAN KEGIATAN", bold=True, underline=True, size=11)

    r2 = t1.rows[2].cells[0]
    p = r2.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_run(p, "{{URAIAN_KEGIATAN}}", size=11)

    r3 = t1.rows[3].cells[0]
    p = r3.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "DISPOSISI DANTONTAR", bold=True, underline=True, size=11)

    r4 = t1.rows[4].cells[0]
    p4 = r4.paragraphs[0]
    p4.paragraph_format.space_before = Pt(24)  # blank space for handwritten disposisi

    new_para(doc, spacing_after=6)

    # ---- signature block ----
    def tabbed(text_left, text_right, bold=False, size=11, tab_at=3.4):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(tab_at))
        add_run(p, text_left, bold=bold, size=size)
        add_run(p, "\t", bold=bold, size=size)
        add_run(p, text_right, bold=bold, size=size)
        return p

    p = new_para(doc, WD_ALIGN_PARAGRAPH.RIGHT)
    add_run(p, "{{TTD_TEMPAT_TANGGAL}}", bold=True, size=11)

    tabbed("{{DANTON_JABATAN}}", "YANG MEMBUAT LAPORAN", bold=True, size=11)
    doc.paragraphs[-1].alignment = None
    new_para(doc, spacing_after=0)
    new_para(doc, spacing_after=0)
    new_para(doc, spacing_after=0)
    tabbed("{{DANTON_NAMA}}", "{{NAMA_TARUNA}}", bold=True, size=11, tab_at=3.4)
    tabbed("{{DANTON_PANGKAT_NRP}}", "{{TARUNA_PANGKAT_NOAK}}", bold=True, size=11, tab_at=3.4)

    new_para(doc, spacing_after=0)
    p = new_para(doc)
    add_run(p, "Mengetahui,", bold=True, size=11)
    p = new_para(doc)
    add_run(p, "{{DANKI_JABATAN}}", bold=True, size=11)
    new_para(doc, spacing_after=0)
    new_para(doc, spacing_after=0)
    p = new_para(doc)
    add_run(p, "{{DANKI_NAMA}}", bold=True, size=11)
    p = new_para(doc)
    add_run(p, "{{DANKI_PANGKAT_NRP}}", bold=True, size=11)

    new_para(doc, spacing_after=6)
    new_para(doc, spacing_after=6)

    p = new_para(doc)
    add_run(p, "LAMPIRAN DOKUMENTASI", bold=True, underline=True, size=12)

    new_para(doc, spacing_after=6)
    p = doc.add_paragraph()
    add_run(p, "{{LAMPIRAN_FOTO}}", size=1)  # marker paragraph; app inserts pictures here

    new_para(doc, spacing_after=12)
    p = new_para(doc, WD_ALIGN_PARAGRAPH.RIGHT)
    add_run(p, "{{LAMP_TEMPAT_TANGGAL}}", bold=True, size=11)
    p = new_para(doc, WD_ALIGN_PARAGRAPH.RIGHT)
    add_run(p, "YANG MEMBUAT LAPORAN", bold=True, size=11)
    new_para(doc, spacing_after=0)
    new_para(doc, spacing_after=0)
    p = new_para(doc, WD_ALIGN_PARAGRAPH.RIGHT)
    add_run(p, "{{NAMA_TARUNA}}", bold=True, size=11)
    p = new_para(doc, WD_ALIGN_PARAGRAPH.RIGHT)
    add_run(p, "{{LAMP_PANGKAT_NOAK}}", bold=True, size=11)

    doc.save('template_lhp.docx')
    print("template_lhp.docx written")


if __name__ == '__main__':
    build()
